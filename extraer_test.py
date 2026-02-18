import docx
import json
import re
import os

def extractor_final(file_path):
    print(f"📂 Procesando archivo: {file_path}")
    try:
        doc = docx.Document(file_path)
        full_text = []

        # 1. Extraer todo el texto limpio (párrafos y tablas)
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if para.text.strip():
                            full_text.append(para.text.strip())
        
        # Unir todo con saltos de línea para procesarlo línea por línea
        texto_completo = "\n".join(full_text)
        lines = texto_completo.split('\n')
        
        print(f"📝 Total líneas de texto extraídas: {len(lines)}")

        # 2. Lógica de Extracción (Versión Flexible)
        preguntas = []
        current_pregunta = None
        
        # Regex flexible: Numero + separador + texto
        re_inicio_pregunta = re.compile(r'^(\d+)[\.\-\)\s]+(.*)')
        
        # Regex para opciones
        re_opcion = re.compile(r'^([a-dA-D])[\)\.\-\s]+(.*)')

        for line in lines:
            line = line.strip()
            if not line: continue

            # Chequear si es Pregunta
            match_preg = re_inicio_pregunta.match(line)
            if match_preg:
                # Guardar la anterior si existe
                if current_pregunta:
                    preguntas.append(current_pregunta)

                num = match_preg.group(1)
                txt = match_preg.group(2).strip()
                
                current_pregunta = {
                    "id": int(num),
                    "tema": "General",
                    "pregunta": txt,
                    "opciones": {},
                    "correcta": "a", 
                    "explicacion": ""
                }
                continue

            # Chequear si es Opción
            if current_pregunta:
                match_opc = re_opcion.match(line)
                if match_opc:
                    letra = match_opc.group(1).lower()
                    texto_opc = match_opc.group(2).strip()
                    current_pregunta["opciones"][letra] = texto_opc
                else:
                    # Continuación de texto
                    if current_pregunta["opciones"]:
                        last_key = list(current_pregunta["opciones"].keys())[-1]
                        current_pregunta["opciones"][last_key] += " " + line
                    else:
                        current_pregunta["pregunta"] += " " + line

        # Añadir la última
        if current_pregunta:
            preguntas.append(current_pregunta)

        # Filtrar preguntas muy vacías (sin texto)
        preguntas = [p for p in preguntas if len(p["pregunta"]) > 5]

        print(f"✅ Preguntas encontradas (sin filtrar opciones): {len(preguntas)}")

        # 3. Guardar JSON
        if preguntas:
            os.makedirs('data', exist_ok=True)
            output_path = 'data/preguntas.json'
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(preguntas, f, ensure_ascii=False, indent=2)
            print(f"💾 Guardado en: {output_path}")

    except Exception as e:
        print(f"❌ Error procesando el archivo: {e}")

if __name__ == "__main__":
    archivo = 'pinche test mad CM.docx'
    extractor_final(archivo)